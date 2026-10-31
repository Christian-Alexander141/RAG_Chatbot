"""Wissensbasis: eigene Dokumente einlesen, in Qdrant indexieren und durchsuchen."""
from __future__ import annotations

import logging
import re
import shutil
import threading
import weakref
from pathlib import Path
from typing import Any, Optional

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from . import app_settings, config, ollama_models
from .bm25 import BM25

logger = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".mdx", ".json", ".csv", ".html", ".xml", ".yml", ".yaml", ".log"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx"}

_client = None
_vector_store = None  # Vector Store des gewählten Embedding-Modells
_store_lock = threading.RLock()
# BM25-Index je Vector Store, wird bei jeder Änderung der Wissensbasis neu aufgebaut
_bm25_indexes: "weakref.WeakKeyDictionary[Any, tuple[BM25, dict[str, Document]]]" = weakref.WeakKeyDictionary()
_bm25_lock = threading.Lock()

# Hybride Suche: so viele Kandidaten liefert jede der beiden Suchen, bevor sie zusammengeführt werden
SEARCH_CANDIDATES = 20
RRF_K = 60  # übliche Konstante der Reciprocal Rank Fusion


class KnowledgeBaseError(RuntimeError):
    pass


def require_ollama_model(model: str) -> None:
    """Prüft, ob das Modell lokal vorhanden ist. Lädt bewusst nichts herunter."""
    import ollama

    client = ollama.Client(host=config.OLLAMA_BASE_URL)
    try:
        client.show(model)
    except ollama.ResponseError as exc:
        raise KnowledgeBaseError(
            f"Das Ollama-Modell '{model}' ist nicht installiert. Installieren mit: ollama pull {model}"
        ) from exc
    except Exception as exc:
        raise KnowledgeBaseError(f"Ollama ist unter {config.OLLAMA_BASE_URL} nicht erreichbar: {exc}") from exc


def same_model(a: str, b: str) -> bool:
    """Ollama ergänzt fehlende Tags um ":latest" ("nomic-embed-text" ist "nomic-embed-text:latest")."""
    def full(name: str) -> str:
        return name if ":" in name else f"{name}:latest"
    return full(a) == full(b)


def embedding_model() -> str:
    """Das Embedding-Modell der Wissensbasis: in den Einstellungen gewählt (gilt für alle), sonst EMBEDDING_MODEL."""
    return app_settings.get("embedding_model") or config.EMBEDDING_MODEL


def collection_name(model: str) -> str:
    """Jedes Embedding-Modell hat eine eigene Collection, denn Vektoren verschiedener Modelle passen nicht zusammen."""
    return f"{config.COLLECTION_NAME}-{re.sub(r'[^a-z0-9]+', '-', model.lower()).strip('-')}"


def _get_client():
    global _client
    if _client is None:
        from qdrant_client import QdrantClient

        config.QDRANT_PATH.mkdir(parents=True, exist_ok=True)
        _client = QdrantClient(path=str(config.QDRANT_PATH))  # der lokale Modus erlaubt nur einen Client je Ordner
    return _client


def _make_embeddings(model: str):
    from langchain_ollama import OllamaEmbeddings

    require_ollama_model(model)
    return OllamaEmbeddings(model=model, base_url=config.OLLAMA_BASE_URL)


def _open_store(model: str):
    from langchain_qdrant import QdrantVectorStore
    from qdrant_client.models import Distance, VectorParams

    embeddings = _make_embeddings(model)
    client = _get_client()
    name = collection_name(model)
    size = len(embeddings.embed_query("dimension probe"))
    if not client.collection_exists(name):
        client.create_collection(name, vectors_config=VectorParams(size=size, distance=Distance.COSINE))
    elif client.get_collection(name).config.params.vectors.size != size:
        raise KnowledgeBaseError(
            f"Die Collection '{name}' passt nicht zu '{model}' (andere Vektorgröße). Bitte {config.QDRANT_PATH} löschen."
        )
    return QdrantVectorStore(client=client, collection_name=name, embedding=embeddings)


def _add_missing_documents(store: Any) -> int:
    """Liest Dokumente aus DOCUMENTS_DIR ein, die in der Collection fehlen: alle bei einem neuen Embedding-Modell,
    sonst die seit dem letzten Wechsel hochgeladenen. Gibt die Anzahl zurück."""
    known = list_sources(store)
    added = 0
    for path in sorted(config.DOCUMENTS_DIR.rglob("*")) if config.DOCUMENTS_DIR.exists() else []:
        source = str(path.relative_to(config.DOCUMENTS_DIR))
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES and source not in known:
            try:
                ingest_file(path, source, vector_store=store)
                added += 1
            except KnowledgeBaseError:
                logger.warning("'%s' konnte nicht eingelesen werden", source, exc_info=True)
    if added:
        logger.info("%d Dokumente in %s eingelesen", added, store.collection_name)
    return added


def _available_embedding_model() -> str:
    """Das gewählte Modell. Wurde keins ausdrücklich gewählt und ist der Standard nicht installiert,
    das erste installierte Embedding-Modell (und diese Wahl merken)."""
    model = embedding_model()
    if app_settings.get("embedding_model"):
        return model
    installed = ollama_models.with_capability("embedding")
    if installed and not any(same_model(model, name) for name in installed):
        logger.warning("Embedding-Modell %s ist nicht installiert, nutze %s", model, installed[0])
        app_settings.put("embedding_model", installed[0])
        return installed[0]
    return model


def get_vector_store():
    global _vector_store
    with _store_lock:
        if _vector_store is None:
            store = _open_store(_available_embedding_model())
            _add_missing_documents(store)
            _vector_store = store
        return _vector_store


def set_embedding_model(model: str) -> int:
    """Wechselt das Embedding-Modell für alle Nutzer. Dokumente, die mit diesem Modell noch nicht eingelesen sind,
    werden jetzt eingelesen (das kann dauern). Gibt deren Anzahl zurück."""
    global _vector_store
    with _store_lock:
        store = _open_store(model)
        added = _add_missing_documents(store)
        app_settings.put("embedding_model", model)
        _vector_store = store
        return added


def load_file(path: Path, source: Optional[str] = None) -> list[Document]:
    """Liest eine Datei und gibt sie als Dokumente zurück (PDF: eine Seite pro Dokument)."""
    source = source or path.name
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader

        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(PdfReader(path).pages)]
        return [
            Document(page_content=text, metadata={"source": source, "page": number})
            for number, text in pages
            if text.strip()
        ]

    if suffix == ".docx":
        from docx import Document as DocxDocument

        text = "\n".join(p.text for p in DocxDocument(path).paragraphs)
    elif suffix in TEXT_SUFFIXES:
        text = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise KnowledgeBaseError(f"Dateityp '{suffix}' wird nicht unterstützt (erlaubt: {sorted(SUPPORTED_SUFFIXES)}).")

    return [Document(page_content=text, metadata={"source": source})] if text.strip() else []


def ingest_file(path: Path, source: Optional[str] = None, vector_store: Any = None) -> int:
    """Indexiert eine Datei neu (vorhandene Chunks derselben Quelle werden ersetzt). Gibt die Chunk-Anzahl zurück."""
    source = source or path.name
    store = vector_store or get_vector_store()
    documents = load_file(path, source)
    if not documents:
        raise KnowledgeBaseError(f"'{source}' enthält keinen lesbaren Text.")

    chunks = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120).split_documents(documents)
    delete_source(source, vector_store=store)
    store.add_documents(chunks)
    _bm25_indexes.pop(store, None)
    return len(chunks)


def add_document(source_path: Path, name: str, vector_store: Any = None) -> int:
    """Legt eine Datei im Dokumentenordner ab und indexiert sie. Bei Fehlern wird nichts zurückgelassen."""
    name = Path(name).name  # nie Pfadanteile aus einem Upload übernehmen
    if Path(name).suffix.lower() not in SUPPORTED_SUFFIXES:
        raise KnowledgeBaseError(f"'{name}': Dateityp nicht unterstützt (erlaubt: {sorted(SUPPORTED_SUFFIXES)}).")
    config.DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    target = config.DOCUMENTS_DIR / name
    if source_path.resolve() != target.resolve():
        shutil.copyfile(source_path, target)
    try:
        return ingest_file(target, name, vector_store=vector_store)
    except Exception:
        target.unlink(missing_ok=True)
        raise


def remove_document(name: str, vector_store: Any = None) -> None:
    """Entfernt ein Dokument, und zwar aus den Collections aller Embedding-Modelle. Sonst tauchte es nach einem
    Wechsel zurück auf ein früheres Modell wieder auf."""
    name = Path(name).name
    if vector_store is not None:
        delete_source(name, vector_store=vector_store)
    else:
        store = get_vector_store()
        delete_source(name, vector_store=store)
        prefix = f"{config.COLLECTION_NAME}-"
        for collection in store.client.get_collections().collections:
            if collection.name.startswith(prefix) and collection.name != store.collection_name:
                store.client.delete(collection_name=collection.name, points_selector=_source_filter(name))
    (config.DOCUMENTS_DIR / name).unlink(missing_ok=True)


def ingest_directory(directory: Path = config.DOCUMENTS_DIR, vector_store: Any = None) -> dict[str, int]:
    """Indexiert alle unterstützten Dateien eines Ordners. Gibt {Quelle: Chunks} zurück."""
    result: dict[str, int] = {}
    if not directory.exists():
        return result
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            source = str(path.relative_to(directory))
            try:
                result[source] = ingest_file(path, source, vector_store=vector_store)
            except KnowledgeBaseError:
                continue
    return result


def _source_filter(source: str):
    from qdrant_client.models import FieldCondition, Filter, MatchValue

    return Filter(must=[FieldCondition(key="metadata.source", match=MatchValue(value=source))])


def delete_source(source: str, vector_store: Any = None) -> None:
    store = vector_store or get_vector_store()
    store.client.delete(collection_name=store.collection_name, points_selector=_source_filter(source))
    _bm25_indexes.pop(store, None)


def list_sources(vector_store: Any = None) -> dict[str, int]:
    """Gibt {Quelle: Anzahl Chunks} aller indexierten Dokumente zurück."""
    store = vector_store or get_vector_store()
    counts: dict[str, int] = {}
    offset = None
    while True:
        points, offset = store.client.scroll(
            collection_name=store.collection_name, limit=256, offset=offset, with_payload=["metadata"]
        )
        for point in points:
            source = (point.payload or {}).get("metadata", {}).get("source")
            if source:
                counts[source] = counts.get(source, 0) + 1
        if offset is None:
            return counts


def _bm25_index(store: Any) -> tuple[BM25, dict[str, Document]]:
    """BM25 über alle Textabschnitte der Wissensbasis. Wird erst bei Bedarf gebaut und dann wiederverwendet."""
    with _bm25_lock:
        if store not in _bm25_indexes:
            chunks: dict[str, Document] = {}
            offset = None
            while True:
                points, offset = store.client.scroll(
                    collection_name=store.collection_name, limit=256, offset=offset, with_payload=True
                )
                for point in points:
                    payload = point.payload or {}
                    metadata = {**payload.get("metadata", {}), "_id": str(point.id)}
                    chunks[str(point.id)] = Document(page_content=payload.get("page_content", ""), metadata=metadata)
                if offset is None:
                    break
            _bm25_indexes[store] = (BM25({i: d.page_content for i, d in chunks.items()}), chunks)
        return _bm25_indexes[store]


def search(query: str, k: int = 4, vector_store: Any = None) -> list[Document]:
    """Hybride Suche: Embeddings (Bedeutung) und BM25 (exakte Begriffe), zusammengeführt per Reciprocal Rank Fusion.
    Mit RETRIEVAL_MODE=dense nur Embeddings, z. B. zum Vergleich."""
    store = vector_store or get_vector_store()
    dense = store.similarity_search(query, k=SEARCH_CANDIDATES if config.RETRIEVAL_MODE == "hybrid" else k)
    if config.RETRIEVAL_MODE != "hybrid":
        return dense[:k]

    bm25, chunks = _bm25_index(store)
    scores: dict[str, float] = {}
    found: dict[str, Document] = {}
    lexical = [(chunk_id, chunks[chunk_id]) for chunk_id, _ in bm25.search(query, SEARCH_CANDIDATES)]
    semantic = [(doc.metadata.get("_id"), doc) for doc in dense]
    for ranking in (semantic, lexical):
        for rank, (chunk_id, doc) in enumerate(ranking):
            scores[chunk_id] = scores.get(chunk_id, 0.0) + 1 / (RRF_K + rank + 1)
            found.setdefault(chunk_id, doc)
    return [found[i] for i in sorted(scores, key=scores.__getitem__, reverse=True)[:k]]
